
import paho.mqtt.client as mqtt
import ssl
import os
from dotenv import load_dotenv
from datetime import datetime
import openpyxl
import smtplib
from email.mime.text import MIMEText

# Charger les variables d'environnement
load_dotenv()

broker = "cac.campiserveur.com"
port = int(os.getenv("MQTT_PORT", 8886))  # valeur par défaut
username = os.getenv("MQTT_USERNAME")
password = os.getenv("MQTT_PASSWORD")
client_cert = "client_certificate_prod.pem"
client_key = "client_prod.pem"
ca_cert = "combined_ca_prod.pem"

#MAIL_TO = os.getenv("ALERT_EMAIL")
#MAIL_FROM = os.getenv("SMTP_SENDER")
#SMTP_SERVER = os.getenv("SMTP_SERVER")
#SMTP_PORT = int(os.getenv("SMTP_PORT", 587))
#SMTP_USER = os.getenv("SMTP_USER")
#SMTP_PASS = os.getenv("SMTP_PASSWORD")

connected_topic = "+/+/connected"
fw_version_topic = "+/+/fw_version"

current_date = datetime.now().strftime("%Y-%m-%d")
excel_file = "alertes_mqtt.xlsx"
"""
def send_email_alert(serial, payload):
    if not MAIL_TO or not SMTP_SERVER:
        print("❌ Configuration email manquante")
        return

    subject = f"⚠️ Alerte version basse - {serial}"
    body = f"L'équipement {serial} a une version de firmware trop basse : {payload}"
    msg = MIMEText(body)
    msg["Subject"] = subject
    msg["From"] = MAIL_FROM or SMTP_USER
    msg["To"] = MAIL_TO

    try:
        with smtplib.SMTP(SMTP_SERVER, SMTP_PORT) as server:
            server.starttls()
            server.login(SMTP_USER, SMTP_PASS)
            server.send_message(msg)
        print(f"📧 Alerte envoyée à {MAIL_TO}")
    except Exception as e:
        print(f"❌ Erreur lors de l'envoi de l'email : {e}")
"""
def get_type_from_serial(serial):
    if serial.startswith("1101-"):
        return "centrale"
    elif len(serial) == 32 and serial.isalnum():
        return "lecteur"
    return "inconnu"

def write_alert(sheetname, data):
    if not os.path.exists(excel_file):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    else:
        wb = openpyxl.load_workbook(excel_file)

    if sheetname not in wb.sheetnames:
        ws = wb.create_sheet(sheetname)
        ws.append(["Horodatage", "Type", "Serial", "Valeur", "Alerte", "Catégorie"])
    else:
        ws = wb[sheetname]

    ws.append(data)
    wb.save(excel_file)

def extract_serial(topic):
    parts = topic.split("/")
    if len(parts) >= 2:
        return parts[1]
    return "inconnu"

def on_message(client, userdata, msg):
    topic = msg.topic
    payload = msg.payload.decode().strip()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    serial = extract_serial(topic)
    category = get_type_from_serial(serial)

    if topic.endswith("connected") and payload == "1":
        userdata[serial] = True
    elif topic.endswith("fw_version"):
        if userdata.get(serial) and payload < "1.3.5":
            print(f"[{now}] ⚠️ ALERTE: {serial} ({category}) version trop basse : {payload}")
            write_alert(current_date, [
                now, "fw_version", serial, payload, "Version < 1.3.5", category
            ])
            #send_email_alert(serial, payload)
        userdata[serial] = False

def main():
    userdata = {}
    client = mqtt.Client(userdata=userdata)
    #client.username_pw_set(username, password)
    context = ssl.create_default_context()
    context.set_ciphers("DEFAULT:@SECLEVEL=1")
    context.load_cert_chain(certfile=client_cert, keyfile=client_key)
    context.load_verify_locations(cafile=ca_cert)
    client.tls_set_context(context)
    client.tls_insecure_set(False)

    client.on_message = on_message

    client.connect(broker, port)
    client.subscribe(connected_topic, qos=1)
    client.subscribe(fw_version_topic, qos=1)

    print("📡 Écoute des topics...")
    client.loop_forever()

if __name__ == "__main__":
    main()
