
import paho.mqtt.client as mqtt
import ssl
import os
from dotenv import load_dotenv
from datetime import datetime
import openpyxl

# Charger les variables d'environnement
load_dotenv()

broker = os.getenv("MQTT_BROKER")
port = int(os.getenv("MQTT_PORT", 8886))
username = os.getenv("MQTT_USERNAME")
password = os.getenv("MQTT_PASSWORD")
client_cert = os.getenv("MQTT_CLIENT_CERT")
client_key = os.getenv("MQTT_CLIENT_KEY")
ca_cert = os.getenv("MQTT_CA_CERT")

serial_target = "22465200002400002b00002508000804"
connected_topic = f"general/{serial_target}/connected"
fw_version_topic = f"general/{serial_target}/fw_version"

current_date = datetime.now().strftime("%Y-%m-%d")
excel_file = "alertes_mqtt.xlsx"

def write_alert(sheetname, data):
    if not os.path.exists(excel_file):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
    else:
        wb = openpyxl.load_workbook(excel_file)

    if sheetname not in wb.sheetnames:
        ws = wb.create_sheet(sheetname)
        ws.append(["Horodatage", "Type", "Serial", "Valeur", "Alerte"])
    else:
        ws = wb[sheetname]

    ws.append(data)
    wb.save(excel_file)

def on_message(client, userdata, msg):
    topic = msg.topic
    payload = msg.payload.decode().strip()
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    if topic.endswith("connected") and payload == "1":
        userdata["connected"] = True
    elif topic.endswith("fw_version"):
        if userdata.get("connected") and payload < "1.3.8":
            print(f"[{now}] ⚠️ ALERTE: Version trop basse: {payload}")
            write_alert(current_date, [now, "fw_version", serial_target, payload, "Version < 1.3.8"])
        userdata["connected"] = False  # Reset for next session

def main():
    userdata = {"connected": False}
    client = mqtt.Client(userdata=userdata)
    client.username_pw_set(username, password)
    context = ssl.create_default_context()
    context.set_ciphers("DEFAULT:@SECLEVEL=1")
    context.load_cert_chain(certfile=client_cert, keyfile=client_key)
    context.load_verify_locations(cafile=ca_cert)
    client.tls_set_context(context)
    client.tls_insecure_set(False)

    client.on_message = on_message

    client.connect(broker, port)
    client.subscribe(connected_topic, qos=1)
    client.subscribe(fw_version_topic, qos=1)

    print("📡 Écoute des topics...")
    client.loop_forever()

if __name__ == "__main__":
    main()
